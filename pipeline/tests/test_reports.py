import json
import shutil

from docx import Document
from PIL import Image

from pipeline import charts
from pipeline.cli import main
from pipeline.config import ROOT
from pipeline.docx_export import convert
from pipeline.reports import build_all

DATA = ROOT / "web" / "public" / "data"


def _with_corpus(project):
    """Проект с готовым корпусом (копия web/public/data) — без пересборки."""
    out = project.path("output")
    out.mkdir(parents=True, exist_ok=True)
    for name in ("corpus.json", "stats.json"):
        shutil.copy(DATA / name, out / name)
    return project


def test_charts_are_300_dpi(tmp_path):
    path = charts.grouped_bars(tmp_path / "a.png", ["x", "y"], {"ряд": [0.5, 0.25]},
                               title="Проверка")
    with Image.open(path) as image:
        dpi = image.info["dpi"]
        assert round(dpi[0]) == 300 and round(dpi[1]) == 300
        # ровно 16 см по ширине — ширина полосы набора A4
        assert abs(image.width / 300 * 2.54 - 16.0) < 0.05


def test_committed_report_charts_are_300_dpi():
    pngs = sorted((ROOT / "reports").glob("*.png"))
    assert pngs, "в reports/ нет графиков"
    for png in pngs:
        with Image.open(png) as image:
            assert round(image.info["dpi"][0]) == 300, png.name


def test_build_all_without_gold_and_tests(project):
    config = _with_corpus(project)
    lines: list[str] = []
    steps = build_all(config, use_llm=False, log=lines.append)
    status = {s.report: s.status for s in steps}
    assert status["corpus.md"] == "готово"
    assert status["evaluation.md"].startswith("нет данных")
    assert status["aligners.md"].startswith("готово без метрик")
    assert status["errors.md"].startswith("нет данных")
    assert status["approbation.md"].startswith("нет данных")
    out = config.path("reports")
    assert sorted(p.name for p in (out / "docx").glob("*.docx")) == [
        "aligners.docx", "approbation.docx", "corpus.docx", "errors.docx", "evaluation.docx"]
    assert not config.path("gold").exists()
    corpus_md = (out / "corpus.md").read_text(encoding="utf-8")
    totals = json.loads((DATA / "stats.json").read_text(encoding="utf-8"))["totals"]
    assert (f"Текстов: **{totals['texts']}**, пар (троек предложений EN–ZH–RU): "
            f"**{totals['pairs']}**") in corpus_md
    assert "## Регистры" in corpus_md and "## Оценка сложности" in corpus_md

    doc = Document(str(out / "docx" / "corpus.docx"))
    assert len(doc.tables) >= 5
    assert len(doc.inline_shapes) == 5  # графики из раздела «Файлы»
    captions = [p.text for p in doc.paragraphs if p.text.startswith(("Таблица", "Рисунок"))]
    assert captions[0] == "Таблица 1 — Тексты"
    assert "Рисунок 2 — Распределение падежей существительных (RU)" in captions
    assert doc.styles["Normal"].font.name == "Times New Roman"


def test_docx_conversion_of_markdown_elements(tmp_path):
    md = tmp_path / "r.md"
    png = charts.grouped_bars(tmp_path / "evaluation_prf.png", ["a"], {"b": [0.5]}, title="t")
    md.write_text("\n".join([
        "# Заголовок", "", "Абзац с **жирным**, *курсивом* и `кодом`.", "",
        "> Предупреждение", "", "- пункт", "  - вложенный", "", "1. первый", "",
        "## Таблица данных", "", "| Имя | Число |", "|---|---:|", "| а \\| б | 1,5 |", "",
        "```", "python -m pipeline reports", "```", "", "## Файлы", "",
        f"- `{png.name}`", "- `data.csv`", "",
    ]), encoding="utf-8")
    doc = Document(str(convert(md, tmp_path / "r.docx")))
    texts = [p.text for p in doc.paragraphs]
    assert "Заголовок" in texts and "Абзац с жирным, курсивом и кодом." in texts
    bold = [r.text for p in doc.paragraphs for r in p.runs if r.bold]
    assert "жирным" in bold
    assert doc.tables[0].cell(1, 0).text == "а | б"
    assert doc.tables[0].cell(1, 1).paragraphs[0].alignment == 2  # по правому краю
    assert "Таблица 1 — Таблица данных" in texts
    assert "Рисунок 1 — Точность, полнота и F1 автоматической разметки по эталону" in texts
    assert "Файлы" not in texts and "python -m pipeline reports" in texts
    styles = {p.style.name for p in doc.paragraphs}
    assert {"List Bullet", "List Bullet 2", "List Number"} <= styles


def test_cli_reports(project, monkeypatch):
    config = _with_corpus(project)
    monkeypatch.setattr("pipeline.cli.load_config", lambda path=None: config)
    assert main(["reports", "--no-llm", "--no-docx"]) == 0
    assert (config.path("reports") / "corpus.md").exists()
    assert not (config.path("reports") / "docx").exists()
